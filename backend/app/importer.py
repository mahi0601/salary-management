"""Reads an uploaded .xlsx or .csv into validated employees. No database in here:
the "email already exists" check is done by the caller."""
import csv
import io
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any, Generator

from openpyxl import load_workbook
from pydantic import ValidationError

from app.errors import ImportFileError
from app.schemas import EmployeeIn

REQUIRED_COLUMNS = ["full_name", "email", "job_title", "department", "country", "salary", "hire_date"]
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20_000

TEMPLATE_CSV = (
    ",".join(REQUIRED_COLUMNS) + "\n"
    "Jane Doe,jane.doe@acme.com,Software Engineer,Engineering,US,120000,2023-05-01\n"
)


@dataclass(frozen=True)
class RowError:
    row: int  # the row number as the person sees it in Excel (header is row 1)
    field: str
    message: str


@dataclass
class ParsedImport:
    total_rows: int
    employees: list[tuple[int, EmployeeIn]]
    errors: list[RowError]


def parse(filename: str, content: bytes) -> ParsedImport:
    if len(content) > MAX_BYTES:
        raise ImportFileError(f"File is larger than {MAX_BYTES // (1024 * 1024)} MB")
    rows = _read_table(filename, content)

    employees, errors, first_row_of_email = [], [], {}
    for number, raw in rows:
        try:
            employee = EmployeeIn(**raw)
        except ValidationError as exc:
            errors += [RowError(number, str(e["loc"][0]) if e["loc"] else "row", e["msg"]) for e in exc.errors()]
            continue
        earlier = first_row_of_email.setdefault(employee.email, number)
        if earlier != number:
            errors.append(RowError(number, "email", f"Same email as row {earlier}"))
            continue
        employees.append((number, employee))
    return ParsedImport(total_rows=len(rows), employees=employees, errors=errors)


def _read_table(filename: str, content: bytes) -> list[tuple[int, dict[str, Any]]]:
    name = filename.lower()
    if name.endswith(".csv"):
        table = _csv_rows(content)
    elif name.endswith(".xlsx"):
        table = _xlsx_rows(content)
    else:
        raise ImportFileError("Upload an .xlsx or .csv file")

    try:
        header_row = next(table, None)
        if header_row is None:
            raise ImportFileError("The file is empty")
        header = [_normalise_header(cell) for cell in header_row]
        missing = [c for c in REQUIRED_COLUMNS if c not in header]
        if missing:
            raise ImportFileError("Missing columns: " + ", ".join(missing))

        rows = []
        for number, cells in enumerate(table, start=2):
            if all(c is None or str(c).strip() == "" for c in cells):
                continue  # Excel often leaves blank rows at the bottom
            if len(rows) >= MAX_ROWS:
                raise ImportFileError(f"More than {MAX_ROWS:,} rows; split the file")
            values = dict(zip(header, cells))
            rows.append((number, {c: _clean(c, values.get(c)) for c in REQUIRED_COLUMNS}))
        return rows
    finally:
        table.close()  # releases the workbook if we stopped early


def _normalise_header(cell: object) -> str:
    return str(cell or "").strip().lower().replace(" ", "_")


def _clean(column: str, value: Any) -> Any:
    if isinstance(value, datetime):  # Excel date cells
        return value.date()
    if isinstance(value, date):
        return value
    if column == "salary":
        if isinstance(value, float) and value.is_integer():
            return int(value)
        if isinstance(value, str):
            return value.replace(",", "").strip()  # "2,400,000" typed as text
    if value is None:
        return ""
    return value if isinstance(value, (int, float)) else str(value).strip()


def _csv_rows(content: bytes) -> Generator[list[Any], None, None]:
    try:
        text = content.decode("utf-8-sig")  # Excel writes a BOM
    except UnicodeDecodeError as exc:
        raise ImportFileError("The file is not UTF-8 text; save it as CSV UTF-8") from exc
    yield from csv.reader(io.StringIO(text))


def _xlsx_rows(content: bytes) -> Generator[list[Any], None, None]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many different things for a bad file
        raise ImportFileError("That is not a valid .xlsx file") from exc
    try:
        for row in workbook.worksheets[0].iter_rows(values_only=True):
            yield list(row)
    finally:
        workbook.close()
